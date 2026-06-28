import openpyxl
from copy import copy
from typing import Dict, List, Any
import re

from excel_inspector import ExcelInspector

ILLEGAL_CHARACTERS_RE = re.compile(r'[\000-\010]|[\013-\014]|[\016-\037]')

def clean_string(val):
    if isinstance(val, str):
        return ILLEGAL_CHARACTERS_RE.sub("", val)
    return val

class ExcelWriter:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.inspector = ExcelInspector(file_path)
        self.wb = openpyxl.load_workbook(file_path)
        self.stats = {
            "DATA": {"written": 0, "start_row": None, "end_row": None},
            "DATA_DA": {"written": 0, "start_row": None, "end_row": None}
        }
        
    def _find_last_data_row(self, sheet) -> int:
        for r in range(sheet.max_row, 0, -1):
            if sheet.cell(row=r, column=1).value is not None or sheet.cell(row=r, column=2).value is not None:
                return r
        return 1

    def _set_cell(self, sheet, r, c, val):
        cell = sheet.cell(row=r, column=c)
        if type(cell).__name__ != 'MergedCell':
            cell.value = clean_string(val)

    def append_data(self, unified_data_list: List[Dict[str, Any]], logs: List[Dict[str, Any]], start_year: int, end_year: int):
        for sheet_key in ["DATA", "DATA_DA"]:
            target_sheet_name = "DATA " if sheet_key == "DATA" else "DATA DA"
            if target_sheet_name not in self.wb.sheetnames:
                if sheet_key == "DATA" and "DATA" in self.wb.sheetnames:
                    target_sheet_name = "DATA"
                elif target_sheet_name not in self.wb.sheetnames:
                    continue
                
            sheet = self.wb[target_sheet_name]
            header_row, headers = self.inspector.find_header_row(target_sheet_name)
            if header_row == -1:
                continue
                
            last_row = self._find_last_data_row(sheet)
            
            ticker_year_to_row = {}
            id_year_to_row = {}
            id_to_tickers = {}
            max_id = 0
            
            ticker_col = headers.get("MÃ CK")
            year_col = headers.get("Năm")
            id_col = headers.get("ID")
            
            if not (ticker_col and year_col and id_col):
                continue
                
            data_sheet = self.inspector.wb[target_sheet_name]
            
            for r in range(header_row + 1, last_row + 1):
                t = str(data_sheet.cell(row=r, column=ticker_col).value or "").strip().upper()
                y = data_sheet.cell(row=r, column=year_col).value
                i = data_sheet.cell(row=r, column=id_col).value
                
                try:
                    if y is not None: y = int(y)
                except: pass
                
                try:
                    if i is not None: i = int(i)
                except: pass
                
                if t and y:
                    ticker_year_to_row[(t, y)] = r
                if i and y:
                    id_year_to_row[(i, y)] = r
                if i is not None and str(i).strip() != "":
                    try:
                        if int(i) > max_id: max_id = int(i)
                    except (ValueError, TypeError):
                        pass
                    if i not in id_to_tickers:
                        id_to_tickers[i] = set()
                    if t:
                        id_to_tickers[i].add(t)
            
            for comp_data in unified_data_list:
                ticker = comp_data["ticker"].upper()
                
                # Determine ID
                ticker_id = None
                for i, ts in id_to_tickers.items():
                    if ticker in ts:
                        ticker_id = i
                        break
                
                if ticker_id is None:
                    empty_ids = [i for i, ts in id_to_tickers.items() if len(ts) == 0]
                    valid_empty_ids = []
                    req_years = list(range(start_year, end_year + 1))
                    
                    for e_id in empty_ids:
                        has_all = True
                        for ry in req_years:
                            if (e_id, ry) not in id_year_to_row:
                                has_all = False
                                break
                        if has_all:
                            valid_empty_ids.append(e_id)
                            
                    if valid_empty_ids:
                        ticker_id = min(valid_empty_ids)
                        id_to_tickers[ticker_id].add(ticker)
                    else:
                        max_id += 1
                        ticker_id = max_id
                        id_to_tickers[ticker_id] = {ticker}

                for year in range(start_year, end_year + 1):
                    years_dict = comp_data["years"].get(year, {})
                    years_dict_prev = comp_data["years"].get(year - 1, {})
                    
                    target_row = None
                    if (ticker, year) in ticker_year_to_row:
                        target_row = ticker_year_to_row[(ticker, year)]
                    elif (ticker_id, year) in id_year_to_row:
                        target_row = id_year_to_row[(ticker_id, year)]
                    
                    if target_row is None:
                        last_row += 1
                        target_row = last_row
                        
                    self._set_cell(sheet, target_row, year_col, year)
                    self._set_cell(sheet, target_row, id_col, ticker_id)
                    if "TÊN CÔNG TY" in headers:
                        if not sheet.cell(row=target_row, column=headers["TÊN CÔNG TY"]).value:
                            self._set_cell(sheet, target_row, headers["TÊN CÔNG TY"], comp_data["company_name"])
                    if "Ngành" in headers:
                        if not sheet.cell(row=target_row, column=headers["Ngành"]).value:
                            self._set_cell(sheet, target_row, headers["Ngành"], comp_data["industry"])
                    elif "NGÀNH" in headers:
                        if not sheet.cell(row=target_row, column=headers["NGÀNH"]).value:
                            self._set_cell(sheet, target_row, headers["NGÀNH"], comp_data["industry"])
                    self._set_cell(sheet, target_row, ticker_col, ticker)
                    
                    # Prepare row data mapping based on sheet
                    row_data_map = {}
                    if sheet_key == "DATA":
                        row_data_map = {
                            "Tổng tài sản đầu kỳ": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kì": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kì (Tổng tài sản t-1)": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kỳ (Tổng tài sản t-1)": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản năm t": years_dict.get("total_assets", {}),
                            "Tổng tài sản": years_dict.get("total_assets", {}),
                            "Nợ phải trả năm t": years_dict.get("liabilities", {}),
                            "Tổng nợ phải trả": years_dict.get("liabilities", {}),
                            "Dòng tiền từ hoạt động kinh doanh": years_dict.get("operating_cash_flow", {}),
                            "Các khoản phải thu ngắn hạn": years_dict.get("short_term_receivables", {}),
                            "Hàng tồn kho": years_dict.get("inventory", {}),
                            "Tổng hàng tồn kho": years_dict.get("inventory", {}),
                            "Doanh thu thuần về bán hàng và cung cấp dịch vụ": years_dict.get("revenue", {}),
                            "Doanh thu thuần": years_dict.get("revenue", {}),
                            "Lợi nhuận sau thuế thu nhập doanh nghiệp": years_dict.get("net_profit_after_tax", {}),
                            "Lợi nhuận sau thuế": years_dict.get("net_profit_after_tax", {}),
                            "Big4": {"value": "NEED_MANUAL_CHECK", "source": "Manual"},
                            "DIGI": {"value": "NEED_MANUAL_CHECK", "source": "Manual"}
                        }
                    else:
                        row_data_map = {
                            "Lợi nhuận sau thuế": years_dict.get("net_profit_after_tax", {}),
                            "Dòng tiền HĐKD": years_dict.get("operating_cash_flow", {}),
                            "Dòng tiền từ hoạt động kinh doanh": years_dict.get("operating_cash_flow", {}),
                            "Tổng tài sản đầu kỳ": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kì": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kì (Tổng tài sản t-1)": years_dict_prev.get("total_assets", {}),
                            "Tổng tài sản đầu kỳ (Tổng tài sản t-1)": years_dict_prev.get("total_assets", {}),
                            "Khoản phải thu ngắn hạn năm t": years_dict.get("short_term_receivables", {}),
                            "Khoản phải thu ngắn hạn năm t-1": years_dict_prev.get("short_term_receivables", {}),
                            "Khoản phải thu ngắn hạn năm t−1": years_dict_prev.get("short_term_receivables", {}), # alternative dash
                            "PPE (Nguyên giá TSCĐ)": years_dict.get("ppe_gross", {}),
                            "Nguyên giá tài sản cố định": years_dict.get("ppe_gross", {}),
                            "Nguyên giá tài sản cố định (PPE)": years_dict.get("ppe_gross", {}),
                            "Nguyên giá Tài sản cố định hữu hình": years_dict.get("ppe_gross", {}),
                            "ROA": years_dict.get("roa", {}),
                            "Doanh thu năm t": years_dict.get("revenue", {}),
                            "Doanh thu t": years_dict.get("revenue", {}),
                            "Doanh thu thuần về bán hàng và cung cấp dịch vụ": years_dict.get("revenue", {}),
                            "Doanh thu năm t-1": years_dict_prev.get("revenue", {}),
                            "Doanh thu t-1": years_dict_prev.get("revenue", {}),
                            "Doanh thu năm t−1": years_dict_prev.get("revenue", {}), # alternative dash
                            "Doanh thu t−1": years_dict_prev.get("revenue", {}),
                            "Doanh thu thuần của năm trước": years_dict_prev.get("revenue", {}),
                            "Lưu chuyển tiền thuần từ hoạt động kinh doanh": years_dict.get("operating_cash_flow", {}),
                            "Tổng tài sản": years_dict.get("total_assets", {}),
                            "Tổng tài sản / TA": years_dict.get("total_assets", {}),
                            "Các khoản phải thu ngắn hạn": years_dict.get("short_term_receivables", {})
                        }
                        
                    for col_name, col_idx in headers.items():
                        if col_name in row_data_map:
                            val = row_data_map[col_name].get("value")
                            if val is not None:
                                self._set_cell(sheet, target_row, col_idx, val)
                            elif col_name not in ["Big4", "DIGI"]:
                                self._set_cell(sheet, target_row, col_idx, "MISSING")
                            elif col_name in ["Big4", "DIGI"]:
                                if not sheet.cell(row=target_row, column=col_idx).value:
                                    self._set_cell(sheet, target_row, col_idx, "NEED_MANUAL_CHECK")
                        elif col_name == "DA" and sheet_key == "DATA":
                            if not sheet.cell(row=target_row, column=col_idx).value:
                                self._set_cell(sheet, target_row, col_idx, "")
                                
                    self.stats[sheet_key]["written"] += 1
                    if self.stats[sheet_key]["start_row"] is None or target_row < self.stats[sheet_key]["start_row"]:
                        self.stats[sheet_key]["start_row"] = target_row
                    if self.stats[sheet_key]["end_row"] is None or target_row > self.stats[sheet_key]["end_row"]:
                        self.stats[sheet_key]["end_row"] = target_row

        if "SOURCE_LOG" not in self.wb.sheetnames:
            log_sheet = self.wb.create_sheet("SOURCE_LOG")
            log_sheet.append(["Mã CK", "Năm", "Sheet", "Cột", "Tên chỉ tiêu trên FireAnt", "Giá trị lấy được", "Trạng thái", "Ghi chú", "Endpoint API"])
        else:
            log_sheet = self.wb["SOURCE_LOG"]
            
        for log in logs:
            log_sheet.append([
                clean_string(log.get("ticker")),
                clean_string(log.get("year")),
                clean_string(log.get("sheet")),
                clean_string(log.get("column")),
                clean_string(log.get("source")),
                clean_string(log.get("value")),
                clean_string(log.get("status")),
                clean_string(log.get("note")),
                clean_string(log.get("endpoint", ""))
            ])
            
    def save(self, output_path: str):
        self.wb.save(output_path)
