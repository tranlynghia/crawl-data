from __future__ import annotations

import argparse
import gzip
import hashlib
import json
import os
import re
import sys
import time
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
import fitz
import pytesseract
import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill



VERIFIED = "Đã xác minh có sở hữu nhà nước"
AUDIT_PROMPT = """Bạn kiểm định tỷ lệ sở hữu Nhà nước trong một trang báo cáo thường niên Việt Nam.
Chỉ dùng nội dung trang được đưa vào. Đọc đúng dòng cổ đông Nhà nước và đúng cột tỷ lệ sở hữu cổ phần/vốn, không lấy tỷ lệ số lượng cổ đông, ngưỡng cổ đông lớn (>5%), tỷ lệ tổ chức hay ngày giao dịch.
Không mặc định cổ đông tổ chức/Tổng công ty là Nhà nước. Chỉ tính phần được nêu rõ là Nhà nước; không tự cộng nhiều tổ chức.
Không lấy tỷ lệ phần vốn do một thành viên HĐQT đại diện làm tổng tỷ lệ Nhà nước của doanh nghiệp.
Ngày dữ liệu phải gắn với bảng/dòng sở hữu và thuộc 15/12–31/12 của năm báo cáo. Không dùng ngày ở footer hoặc năm trên bìa.
Nếu con số đầu vào sai nhưng bảng ghi rõ con số đúng, trả CORRECT. Nếu đầu vào đúng, trả VALID. Nếu không thể xác định chắc chắn, trả UNCLEAR.
Chỉ trả JSON object: decision (VALID|CORRECT|UNCLEAR), percentage (số hoặc null), data_date (YYYY-MM-DD hoặc null), support (trích đoạn rất ngắn từ dòng chứng minh), reason (ngắn)."""
FALLBACK_TERMS = [
    ("cổ đông nhà nước", 140),
    ("state shareholder", 140),
    ("state shareholders", 140),
    ("cơ cấu cổ đông", 125),
    ("shareholder structure", 125),
    ("shareholding structure", 125),
    ("cơ cấu sở hữu", 120),
    ("ownership structure", 120),
    ("cơ cấu vốn cổ phần", 115),
    ("số lượng cổ đông", 110),
    ("number of shareholders", 110),
    ("danh sách cổ đông lớn", 105),
    ("thông tin cổ đông lớn", 105),
    ("major shareholders", 105),
    ("chi tiết vốn góp của chủ sở hữu", 100),
    ("owners capital contribution", 100),
    ("shareholders", 35),
    ("cổ đông", 35),
]


def normalize(value: str) -> str:
    value = value.casefold().replace("đ", "d")
    value = unicodedata.normalize("NFD", value)
    value = "".join(char for char in value if unicodedata.category(char) != "Mn")
    return re.sub(r"[^a-z0-9%]+", " ", value).strip()


NORMALIZED_TERMS = [(term, normalize(term), weight) for term, weight in FALLBACK_TERMS]


def validated_ratio(review: dict[str, Any], row: dict[str, Any]) -> tuple[float | None, str]:
    parsed = review.get("parsed") or {}
    decision = parsed.get("decision")
    if decision not in {"VALID", "CORRECT"}:
        return None, str(parsed.get("reason") or "Chưa đủ căn cứ")
    try:
        percentage = float(parsed["percentage"])
        data_date = date.fromisoformat(str(parsed["data_date"]))
    except (KeyError, TypeError, ValueError):
        return None, "API thiếu tỷ lệ hoặc ngày hợp lệ"
    if not (0 < percentage <= 100 and data_date.year == int(row["Năm báo cáo"]) and data_date.month == 12 and data_date.day >= 15):
        return None, "Tỷ lệ/ngày ngoài phạm vi"
    reason = str(parsed.get("reason") or "")
    for raw_date in re.findall(r"(?:bảng ngày|danh sách chốt ngày)\s*(\d{1,2}/\d{1,2}/20\d{2})", reason, flags=re.IGNORECASE):
        day, month, year = map(int, raw_date.split("/"))
        if (year, month, day) != (data_date.year, data_date.month, data_date.day):
            return None, "Ngày của bảng/danh sách mâu thuẫn với ngày API chọn"
    evidence = str(row.get("Đoạn/bảng bằng chứng") or "")
    values = [float(raw.replace(",", ".")) for raw in re.findall(r"(?<!\d)(\d{1,3}(?:[.,]\d{1,4})?)\s*%", evidence)]
    has_value = any(abs(value - percentage) < 0.0001 for value in values)
    if not has_value:
        normalized = normalize(evidence)
        if "ty le" in normalized or "percentage" in normalized or "%" in evidence:
            decimals = re.findall(r"(?<!\d)(\d{1,3}[.,]\d{1,4})(?!\d)", evidence)
            has_value = any(abs(float(raw.replace(",", ".")) - percentage) < 0.0001 for raw in decimals)
    if not has_value:
        return None, "Tỷ lệ API không xuất hiện trong đoạn PDF"
    input_percentage = row.get("Tỷ lệ sở hữu Nhà nước (%)")
    if decision == "VALID" and input_percentage is not None and pd.notna(input_percentage):
        if abs(percentage - float(input_percentage)) >= 0.0001:
            return None, "API ghi VALID nhưng đổi tỷ lệ"
    return percentage, reason


def request_review(api_key: str, model: str, row: dict[str, Any], timeout: int) -> dict[str, Any]:
    evidence = str(row.get("Đoạn/bảng bằng chứng") or "")[:6500]
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": AUDIT_PROMPT},
            {"role": "user", "content": json.dumps({
                "ticker": row["Mã CK"],
                "report_year": int(row["Năm báo cáo"]),
                "input_percentage": row.get("Tỷ lệ sở hữu Nhà nước (%)"),
                "input_data_date": row.get("Ngày số liệu"),
                "page_text": evidence,
            }, ensure_ascii=False)},
        ],
        "temperature": 0,
    }
    for attempt in range(3):
        try:
            response = requests.post(
                "https://motapis.com/v1/chat/completions",
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
                json=payload,
                timeout=timeout,
            )
            if response.status_code == 429:
                time.sleep(2 ** (attempt + 1))
                continue
            response.raise_for_status()
            body = response.json()
            content = body["choices"][0]["message"]["content"]
            start, end = content.find("{"), content.rfind("}")
            if start < 0 or end < start:
                raise ValueError("API không trả JSON object")
            return {"parsed": json.loads(content[start : end + 1]), "usage": body.get("usage")}
        except (requests.RequestException, KeyError, ValueError) as exc:
            if attempt == 2:
                raise RuntimeError(f"{type(exc).__name__}: {exc}") from exc
            time.sleep(2 ** attempt)
    raise RuntimeError("API thất bại")


def page_score(text: str, report_year: int) -> tuple[int, str | None]:
    normalized = normalize(text)
    matches = [(term, weight) for term, normalized_term, weight in NORMALIZED_TERMS if normalized_term in normalized]
    if not matches:
        return 0, None
    term, score = max(matches, key=lambda item: item[1])
    if "%" in text:
        score += 15
    if str(report_year) in text:
        score += 8
    if re.search(r"(?:1[5-9]|2\d|3[01])\s*[/.-]\s*12\s*[/.-]\s*20\d{2}", text):
        score += 20
    if any(marker in normalized for marker in ("ty le so huu", "percentage", "so co phieu", "number of shares")):
        score += 12
    return score, term


def excerpt(text: str, term: str | None, limit: int = 900) -> str:
    compact_text = re.sub(r"\s+", " ", text).strip()
    if not compact_text:
        return ""
    if not term:
        return compact_text[:limit]
    normalized_text = normalize(compact_text)
    normalized_term = normalize(term)
    index = normalized_text.find(normalized_term)
    if index < 0:
        return compact_text[:limit]
    ratio = index / max(len(normalized_text), 1)
    center = int(ratio * len(compact_text))
    start = max(0, center - limit // 3)
    return compact_text[start : start + limit]


def load_pages(cache_path: Path, pdf_paths: list[Path] | None = None) -> list[dict[str, Any]]:
    signature = None
    if pdf_paths:
        digest = hashlib.sha256()
        for pdf_path in pdf_paths:
            stat = pdf_path.stat()
            digest.update(f"{pdf_path.name}|{stat.st_size}|{stat.st_mtime_ns}".encode("utf-8"))
        signature = digest.hexdigest()
    if cache_path.exists():
        with gzip.open(cache_path, "rt", encoding="utf-8") as handle:
            payload = json.load(handle)
        if signature is None or payload.get("signature") == signature:
            return payload.get("pages", [])
    if not pdf_paths:
        return []
    pages = []
    for pdf_path in pdf_paths:
        with fitz.open(pdf_path) as document:
            for page_number, page in enumerate(document, start=1):
                page_text = page.get_text("text", sort=True).strip()
                visible_chars = len(re.sub(r"\s", "", page_text))
                pages.append({
                    "pdf_file": pdf_path.name,
                    "page_number": page_number,
                    "text": page_text,
                    "visible_chars": visible_chars,
                    "needs_ocr": visible_chars < 80,
                })
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(cache_path, "wt", encoding="utf-8", compresslevel=4) as handle:
        json.dump({"signature": signature, "pages": pages}, handle, ensure_ascii=False)
    return pages


def ocr_missing_pages(manifest: pd.DataFrame, cache_root: Path, ocr_cache_root: Path, pdf_root: Path, tessdata: Path, max_pages: int, dpi: int) -> dict[str, int]:
    from io import BytesIO

    from PIL import Image

    if tessdata.is_dir():
        os.environ["TESSDATA_PREFIX"] = str(tessdata.resolve())
    ocr_cache_root.mkdir(parents=True, exist_ok=True)
    reports = 0
    page_count = 0
    for item in manifest.to_dict("records"):
        ticker = str(item["Mã CK"]).strip().upper()
        year = int(item["Năm"])
        pdf_paths = sorted((pdf_root / ticker).glob(f"{ticker}_{year}*.pdf"))
        if not pdf_paths:
            continue
        cache_path = cache_root / f"{ticker}_{year}.json.gz"
        pages = load_pages(cache_path, pdf_paths)
        if find_fallback_page(pages, year):
            continue
        output_cache = ocr_cache_root / f"{ticker}_{year}.json.gz"
        if output_cache.exists():
            continue
        by_file: dict[str, list[dict[str, Any]]] = {}
        for page in pages:
            if page.get("needs_ocr"):
                by_file.setdefault(str(page["pdf_file"]), []).append(page)
        selected = []
        for file_pages in by_file.values():
            if len(file_pages) <= max_pages:
                selected.extend(file_pages)
                continue
            indexes = {round(index * (len(file_pages) - 1) / max(max_pages - 1, 1)) for index in range(max_pages)}
            selected.extend(file_pages[index] for index in sorted(indexes))
        if not selected:
            continue
        path_map = {path.name: path for path in pdf_paths}
        ocr_pages = []
        for selected_page in selected:
            with fitz.open(path_map[selected_page["pdf_file"]]) as document:
                page = document[selected_page["page_number"] - 1]
                pixmap = page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), colorspace=fitz.csRGB, alpha=False)
                image = Image.open(BytesIO(pixmap.tobytes("png")))
                page_text = pytesseract.image_to_string(image, lang="vie+eng", config="--psm 6").strip()
            ocr_pages.append({**selected_page, "text": page_text, "visible_chars": len(re.sub(r"\s", "", page_text)), "needs_ocr": False})
        candidates = extract_ocr_candidates(ticker, year, ocr_pages)
        payload = {"ticker": ticker, "year": year, "ocr_pages": len(ocr_pages), "pages": ocr_pages, "candidates": candidates, "errors": []}
        with gzip.open(output_cache, "wt", encoding="utf-8", compresslevel=4) as handle:
            json.dump(payload, handle, ensure_ascii=False)
        reports += 1
        page_count += len(ocr_pages)
        print(f"OCR_PROGRESS reports={reports} pages={page_count} ticker={ticker} year={year}", flush=True)
    return {"Báo cáo OCR mới": reports, "Trang OCR mới": page_count}


def extract_ocr_candidates(ticker: str, year: int, pages: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates = []
    for page in pages:
        text = str(page.get("text") or "")
        dates = re.findall(r"(?<!\d)(\d{1,2})\s*[/.-]\s*(\d{1,2})\s*[/.-]\s*(20\d{2})(?!\d)", text)
        matching_dates = [f"{year}-12-{int(day):02d}" for day, month, found_year in dates if int(found_year) == year and int(month) == 12 and 15 <= int(day) <= 31]
        if not matching_dates:
            continue
        for line in text.splitlines():
            normalized_line = normalize(line)
            if not any(term in normalized_line for term in ("co dong nha nuoc", "so huu nha nuoc", "von nha nuoc", "state shareholder", "state ownership")):
                continue
            values = re.findall(r"(?<!\d)(\d{1,3}(?:[.,]\d{1,4})?)\s*%", line)
            if len(values) != 1:
                continue
            percentage = float(values[0].replace(",", "."))
            if not 0 < percentage <= 100:
                continue
            candidates.append({
                "ticker": ticker,
                "report_year": year,
                "percentage": percentage,
                "data_date": matching_dates[0],
                "pdf_file": page["pdf_file"],
                "page_number": page["page_number"],
                "keyword": "sở hữu Nhà nước",
                "evidence": text[:6500],
                "explicit_state_term": True,
                "score": 1,
            })
    return candidates


def find_fallback_page(pages: list[dict[str, Any]], report_year: int) -> dict[str, Any] | None:
    best = None
    for page in pages:
        text = str(page.get("text") or "")
        if int(page.get("visible_chars") or 0) < 80:
            continue
        score, term = page_score(text, report_year)
        if score <= 0:
            continue
        candidate = {
            "score": score,
            "term": term,
            "pdf_file": page.get("pdf_file"),
            "page_number": int(page.get("page_number") or 0),
            "excerpt": excerpt(text, term),
        }
        if best is None or (candidate["score"], -candidate["page_number"]) > (best["score"], -best["page_number"]):
            best = candidate
    return best


def load_ocr(cache_path: Path) -> dict[str, Any]:
    if not cache_path.exists():
        return {"pages": [], "candidates": [], "ocr_pages": 0, "errors": []}
    with gzip.open(cache_path, "rt", encoding="utf-8") as handle:
        return json.load(handle)


def load_audit(cache_path: Path) -> dict[str, Any]:
    if not cache_path.exists():
        return {"parsed": {"decision": "UNCLEAR", "reason": "Chưa kiểm định lại tỷ lệ cũ"}}
    return json.loads(cache_path.read_text(encoding="utf-8"))


def eligible_ocr_candidate(candidate: dict[str, Any], report_year: int) -> bool:
    try:
        data_date = date.fromisoformat(str(candidate["data_date"]))
        ratio = float(candidate["percentage"])
    except (KeyError, TypeError, ValueError):
        return False
    return (
        bool(candidate.get("explicit_state_term"))
        and 0 < ratio <= 100
        and data_date.year == report_year
        and data_date.month == 12
        and data_date.day >= 15
    )


def audit_candidates(manifest: pd.DataFrame, details: pd.DataFrame, cache_root: Path, ocr_cache_root: Path, audit_cache_root: Path, fallback_audit_root: Path, ocr_audit_root: Path, pdf_root: Path, model: str) -> dict[str, int]:
    api_key = os.environ.get("MOTAPIS_API_KEY")
    if not api_key:
        raise RuntimeError("Thiếu MOTAPIS_API_KEY để dùng --audit")
    for directory in (audit_cache_root, fallback_audit_root, ocr_audit_root):
        directory.mkdir(parents=True, exist_ok=True)
    counts = {"Kiểm định lớp chữ mới": 0, "Kiểm định trang bỏ sót mới": 0, "Kiểm định OCR mới": 0}
    details_by_key = {
        (str(row["Mã CK"]).strip().upper(), int(row["Năm báo cáo"])): row
        for row in details.to_dict("records")
    }
    for item in manifest.to_dict("records"):
        ticker = str(item["Mã CK"]).strip().upper()
        year = int(item["Năm"])
        key = (ticker, year)
        detail = details_by_key.get(key)
        pdf_paths = sorted((pdf_root / ticker).glob(f"{ticker}_{year}*.pdf"))
        if not pdf_paths:
            continue
        if detail and detail.get("Trạng thái") == VERIFIED and pd.notna(detail.get("Tỷ lệ sở hữu Nhà nước (%)")):
            output = audit_cache_root / f"{ticker}_{year}.json"
            if not output.exists():
                review = request_review(api_key, model, detail, 180)
                output.write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8")
                counts["Kiểm định lớp chữ mới"] += 1
            continue
        pages = load_pages(cache_root / f"{ticker}_{year}.json.gz", pdf_paths)
        fallback = find_fallback_page(pages, year)
        if fallback:
            page = next((value for value in pages if value["pdf_file"] == fallback["pdf_file"] and value["page_number"] == fallback["page_number"]), None)
            if page:
                page_text = str(page["text"])
                normalized = normalize(page_text)
                explicit = any(term in normalized for term in ("co dong nha nuoc", "so huu nha nuoc", "von nha nuoc", "state shareholder", "state ownership"))
                late_date = bool(re.search(r"(?<!\d)(?:1[5-9]|2\d|3[01])\s*[/.-]\s*12\s*[/.-]\s*" + str(year) + r"(?!\d)", page_text))
                output = fallback_audit_root / f"{ticker}_{year}.json"
                if explicit and late_date and not output.exists():
                    row = {
                        "Mã CK": ticker,
                        "Năm báo cáo": year,
                        "Tỷ lệ sở hữu Nhà nước (%)": None,
                        "Ngày số liệu": None,
                        "Đoạn/bảng bằng chứng": page_text[:6500],
                    }
                    review = request_review(api_key, model, row, 180)
                    review.update({"pdf_file": fallback["pdf_file"], "page_number": fallback["page_number"], "evidence": row["Đoạn/bảng bằng chứng"]})
                    output.write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8")
                    counts["Kiểm định trang bỏ sót mới"] += 1
        ocr_payload = load_ocr(ocr_cache_root / f"{ticker}_{year}.json.gz")
        candidates = [candidate for candidate in ocr_payload.get("candidates", []) if eligible_ocr_candidate(candidate, year)]
        if candidates:
            candidate = max(candidates, key=lambda value: int(value.get("score") or 0))
            output = ocr_audit_root / f"{ticker}_{year}.json"
            if not output.exists():
                row = {
                    "Mã CK": ticker,
                    "Năm báo cáo": year,
                    "Tỷ lệ sở hữu Nhà nước (%)": candidate["percentage"],
                    "Ngày số liệu": candidate["data_date"],
                    "Đoạn/bảng bằng chứng": candidate["evidence"],
                }
                review = request_review(api_key, model, row, 180)
                review.update({"pdf_file": candidate["pdf_file"], "page_number": candidate["page_number"], "evidence": candidate["evidence"]})
                output.write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding="utf-8")
                counts["Kiểm định OCR mới"] += 1
    return counts


def build_review(manifest_path: Path, result_path: Path, cache_root: Path, ocr_cache_root: Path, audit_cache_root: Path, fallback_audit_root: Path, ocr_audit_root: Path, pdf_root: Path, output_path: Path) -> dict[str, int]:
    if manifest_path.suffix.lower() == ".csv":
        manifest = pd.read_csv(manifest_path, encoding="utf-8-sig")
    else:
        manifest = pd.read_excel(manifest_path)
    if not result_path.exists():
        details = pd.DataFrame()
    elif result_path.suffix.lower() == ".csv":
        details = pd.read_csv(result_path, encoding="utf-8-sig")
    else:
        details = pd.read_excel(result_path, sheet_name="Chi_tiet")
    detail_by_key = {
        (str(row["Mã CK"]).strip().upper(), int(row["Năm báo cáo"])): row
        for row in details.to_dict("records")
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Kiem_tra"
    sheet.append(["Mã CK", "Năm", "Tỷ lệ sở hữu Nhà nước (%)", "Trang PDF"])
    source_sheet = workbook.create_sheet("Nguon_kiem_tra")
    source_sheet.append(["Mã CK", "Năm", "File PDF", "Link nguồn", "Ngày số liệu", "Phương pháp", "Trạng thái", "Từ khóa", "Đoạn bằng chứng/trang", "Số trang OCR"])
    notes_sheet = workbook.create_sheet("Ghi_chu")
    notes_sheet.append(["Mục", "Giải thích"])
    notes_sheet.append(["Phạm vi", "101 mã CK × 6 năm = 606 dòng; 605 báo cáo đã có PDF, thiếu CHP 2020."])
    notes_sheet.append(["Cuối năm", "Chỉ lấy ngày dữ liệu 15/12–31/12 của năm báo cáo, không mặc định cả năm."])
    notes_sheet.append(["Không có", "Chưa tìm được tỷ lệ đủ điều kiện từ tài liệu đã đọc; KHÔNG có nghĩa tỷ lệ Nhà nước bằng 0%."])
    notes_sheet.append(["OCR", "OCR lấy mẫu trang scan để tìm ứng viên, chưa đọc hết mọi trang scan. Ứng viên OCR phải kiểm tra tay."])
    notes_sheet.append(["Kiểm định", "Tỷ lệ lớp chữ qua kiểm định API vẫn cần người dùng đối chiếu đúng dòng/cột trên trang PDF."])

    statistics = {
        "Tổng dòng mã-năm": 0,
        "Đã tìm kiếm PDF": 0,
        "Thiếu PDF": 0,
        "Tỷ lệ lớp chữ qua kiểm định": 0,
        "Tỷ lệ bổ sung từ trang bỏ sót": 0,
        "Tỷ lệ OCR qua kiểm định": 0,
        "Không có tỷ lệ nhưng có trang cổ đông": 0,
        "Không tìm thấy trang cổ đông bằng lớp chữ": 0,
        "Số trang đã OCR": 0,
    }

    rows = manifest.sort_values(["Mã CK", "Năm"]).to_dict("records")
    for excel_row, item in enumerate(rows, start=2):
        ticker = str(item["Mã CK"]).strip().upper()
        year = int(item["Năm"])
        detail = detail_by_key.get((ticker, year), {})
        cache_path = cache_root / f"{ticker}_{year}.json.gz"
        pdf_paths = sorted((pdf_root / ticker).glob(f"{ticker}_{year}*.pdf"))
        ocr_payload = load_ocr(ocr_cache_root / f"{ticker}_{year}.json.gz")
        statistics["Số trang đã OCR"] += int(ocr_payload.get("ocr_pages") or 0)
        has_pdf = bool(pdf_paths)
        status = detail.get("Trạng thái")
        old_ratio_present = status == VERIFIED and pd.notna(detail.get("Tỷ lệ sở hữu Nhà nước (%)"))
        audit_ratio = None
        audit_reason = ""
        audit_data_date = None
        if old_ratio_present:
            review = load_audit(audit_cache_root / f"{ticker}_{year}.json")
            audit_ratio, audit_reason = validated_ratio(review, detail)
            audit_data_date = (review.get("parsed") or {}).get("data_date")
        verified = old_ratio_present and audit_ratio is not None

        statistics["Tổng dòng mã-năm"] += 1
        statistics["Đã tìm kiếm PDF"] += int(has_pdf)
        statistics["Thiếu PDF"] += int(not has_pdf)

        pdf_file = None
        page_number: int | None = None
        page_comment = ""
        data_date = None
        keyword = None
        method = None
        evidence = None
        result_status = "Không tìm thấy tỷ lệ đủ điều kiện"
        if verified:
            ratio_value: float | str = float(audit_ratio)
            page_number = int(detail["Trang PDF"])
            pdf_file = str(detail["File PDF"])
            page_comment = (
                f"Tỷ lệ lớp chữ đã qua kiểm định API, vẫn cần kiểm tra tay.\nFile: {pdf_file}\nNgày số liệu: {audit_data_date}\n"
                f"Lý do kiểm định: {audit_reason}\n"
                f"Đoạn kiểm tra: {str(detail.get('Đoạn/bảng bằng chứng') or '')[:1200]}"
            )
            data_date = audit_data_date
            keyword = detail.get("Từ khóa tìm thấy")
            method = f"{detail.get('Phương pháp trích')} + kiểm định Motapis"
            evidence = detail.get("Đoạn/bảng bằng chứng")
            result_status = "Có tỷ lệ lớp chữ qua kiểm định; cần kiểm tra tay"
            statistics["Tỷ lệ lớp chữ qua kiểm định"] += 1
        else:
            ratio_value = "Không có"
            fallback_review_path = fallback_audit_root / f"{ticker}_{year}.json"
            fallback_ratio = None
            fallback_review = None
            if fallback_review_path.exists():
                fallback_review = json.loads(fallback_review_path.read_text(encoding="utf-8"))
                fallback_row = {
                    "Năm báo cáo": year,
                    "Tỷ lệ sở hữu Nhà nước (%)": None,
                    "Đoạn/bảng bằng chứng": fallback_review.get("evidence"),
                }
                fallback_ratio, _ = validated_ratio(fallback_review, fallback_row)
            if fallback_ratio is not None and fallback_review:
                ratio_value = fallback_ratio
                pdf_file = str(fallback_review["pdf_file"])
                page_number = int(fallback_review["page_number"])
                data_date = fallback_review["parsed"].get("data_date")
                keyword = "Cổ đông Nhà nước"
                method = "PyMuPDF + kiểm định Motapis"
                evidence = fallback_review.get("evidence")
                result_status = "Tỷ lệ bổ sung qua kiểm định; cần kiểm tra tay"
                page_comment = (
                    "Tỷ lệ bổ sung từ trang lớp chữ; cần kiểm tra trực tiếp dòng/cột/ngày.\n"
                    f"File: {pdf_file}\nNgày số liệu: {data_date}\n"
                    f"Lý do: {fallback_review['parsed'].get('reason')}\n"
                    f"Đoạn kiểm tra: {str(evidence or '')[:1000]}"
                )
                statistics["Tỷ lệ bổ sung từ trang bỏ sót"] += 1
            else:
                candidates = [candidate for candidate in ocr_payload.get("candidates", []) if eligible_ocr_candidate(candidate, year)]
                best_candidate = max(candidates, key=lambda candidate: int(candidate.get("score") or 0), default=None)
                ocr_ratio = None
                ocr_reason = ""
                ocr_review_path = ocr_audit_root / f"{ticker}_{year}.json"
                if best_candidate and ocr_review_path.exists():
                    ocr_review = json.loads(ocr_review_path.read_text(encoding="utf-8"))
                    ocr_ratio, ocr_reason = validated_ratio(ocr_review, {
                        "Năm báo cáo": year,
                        "Tỷ lệ sở hữu Nhà nước (%)": best_candidate["percentage"],
                        "Đoạn/bảng bằng chứng": best_candidate["evidence"],
                    })
                if best_candidate and ocr_ratio is not None:
                    ratio_value = float(ocr_ratio)
                    pdf_file = str(best_candidate["pdf_file"])
                    page_number = int(best_candidate["page_number"])
                    data_date = (ocr_review.get("parsed") or {}).get("data_date")
                    keyword = best_candidate.get("keyword")
                    method = "Tesseract OCR vie+eng"
                    evidence = best_candidate.get("evidence")
                    result_status = "Tỷ lệ OCR qua kiểm định; cần kiểm tra tay"
                    page_comment = (
                        "Tỷ lệ OCR qua kiểm định API, vẫn cần đối chiếu trực tiếp trên PDF.\n"
                        f"File: {pdf_file}\nNgày số liệu: {data_date}\n"
                        f"Lý do: {ocr_reason}\n"
                        f"Đoạn kiểm tra: {str(evidence or '')[:1200]}"
                    )
                    statistics["Tỷ lệ OCR qua kiểm định"] += 1
                else:
                    text_pages = load_pages(cache_path, pdf_paths)
                    fallback = find_fallback_page(text_pages + ocr_payload.get("pages", []), year)
                    if best_candidate and ocr_ratio is None:
                        fallback = {
                            "pdf_file": best_candidate["pdf_file"],
                            "page_number": int(best_candidate["page_number"]),
                            "term": "Ứng viên OCR chưa qua kiểm định",
                            "score": 0,
                            "excerpt": f"{ocr_reason}; {str(best_candidate.get('evidence') or '')[:800]}",
                        }
                    if old_ratio_present and audit_ratio is None and detail.get("File PDF") and pd.notna(detail.get("Trang PDF")):
                        fallback = {
                            "pdf_file": detail["File PDF"],
                            "page_number": int(detail["Trang PDF"]),
                            "term": "Tỷ lệ cũ không qua kiểm định",
                            "score": 0,
                            "excerpt": f"{audit_reason}; {str(detail.get('Đoạn/bảng bằng chứng') or '')[:800]}",
                        }
                    if fallback:
                        pdf_file = str(fallback["pdf_file"])
                        page_number = int(fallback["page_number"])
                        keyword = fallback["term"]
                        method = "PyMuPDF/Tesseract OCR + từ khóa"
                        evidence = fallback["excerpt"]
                        result_status = "Không có tỷ lệ; có trang cổ đông để kiểm tra"
                        page_comment = (
                            "Chưa có tỷ lệ đủ điều kiện; đây là trang cổ đông để kiểm tra tay.\n"
                            f"File: {pdf_file}\nTừ khóa: {fallback['term']}\nĐiểm: {fallback['score']}\n"
                            f"Đoạn kiểm tra: {fallback['excerpt']}"
                        )
                        statistics["Không có tỷ lệ nhưng có trang cổ đông"] += 1
                    elif has_pdf:
                        result_status = "Chưa tìm thấy trang cổ đông; cần OCR thêm/kiểm tra tay"
                        statistics["Không tìm thấy trang cổ đông bằng lớp chữ"] += 1
                    else:
                        result_status = "Thiếu PDF"

        page_value: int | str
        if page_number:
            page_value = page_number
        elif has_pdf:
            page_value = "Không tìm thấy"
        else:
            page_value = "Không có PDF"
        sheet.append([ticker, year, ratio_value, page_value])
        source_sheet.append([ticker, year, pdf_file, item.get("URL"), data_date, method, result_status, keyword, str(evidence or "")[:3000], int(ocr_payload.get("ocr_pages") or 0)])

        page_cell = sheet.cell(excel_row, 4)
        if pdf_file and page_number:
            pdf_path = (pdf_root / ticker / pdf_file).resolve()
            if pdf_path.exists():
                page_cell.hyperlink = f"{pdf_path.as_uri()}#page={page_number}"
                page_cell.style = "Hyperlink"
            page_cell.comment = Comment(page_comment, "Codex")
            if result_status.startswith("Ứng viên OCR"):
                sheet.cell(excel_row, 3).comment = Comment("Tỷ lệ OCR là ứng viên, chưa xác minh độc lập. Kiểm tra trực tiếp trang PDF.", "Codex")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 12, "B": 10, "C": 30, "D": 18}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = "0"
        if isinstance(row[2].value, (int, float)):
            row[2].number_format = "0.0000"
        if isinstance(row[3].value, (int, float)):
            row[3].number_format = "0"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in source_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    source_sheet.freeze_panes = "A2"
    source_sheet.auto_filter.ref = source_sheet.dimensions
    for column, width in {"A": 12, "B": 10, "C": 28, "D": 65, "E": 16, "F": 28, "G": 52, "H": 24, "I": 90, "J": 14}.items():
        source_sheet.column_dimensions[column].width = width
    for cell in notes_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    notes_sheet.column_dimensions["A"].width = 22
    notes_sheet.column_dimensions["B"].width = 105
    for row in notes_sheet.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    except PermissionError as exc:
        raise RuntimeError(f"Không thể ghi {output_path}; hãy đóng file Excel đang mở hoặc dùng --output với tên khác") from exc
    return statistics


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Xử lý PDF sở hữu Nhà nước và tạo bảng 606 mã-năm")
    parser.add_argument("--manifest", type=Path, default=Path("danh_muc_pdf_101x6.csv"))
    parser.add_argument("--results", type=Path, default=Path("ownership_extract_101x6/ratio_candidates.csv"))
    parser.add_argument("--cache-root", type=Path, default=Path("ownership_extract_101x6/page_cache"))
    parser.add_argument("--pdf-root", type=Path, default=Path("pdf_101x6"))
    parser.add_argument("--ocr-cache-root", type=Path, default=Path("ownership_extract_101x6/ocr_cache_v2_160_6"))
    parser.add_argument("--audit-cache-root", type=Path, default=Path("ownership_extract_101x6/audit_text_ratios"))
    parser.add_argument("--fallback-audit-root", type=Path, default=Path("ownership_extract_101x6/audit_fallback_ratios"))
    parser.add_argument("--ocr-audit-root", type=Path, default=Path("ownership_extract_101x6/audit_ocr_ratios"))
    parser.add_argument("--ocr-missing", action="store_true", help="OCR lấy mẫu báo cáo chưa tìm thấy trang cổ đông")
    parser.add_argument("--audit", action="store_true", help="Kiểm định ứng viên mới bằng Motapis; cần MOTAPIS_API_KEY")
    parser.add_argument("--max-pages", type=int, default=6)
    parser.add_argument("--dpi", type=int, default=160)
    parser.add_argument("--tessdata", type=Path, default=Path("ownership_cli/tessdata"))
    parser.add_argument("--audit-model", default="gpt-5.6-sol")
    parser.add_argument("--output", type=Path, default=Path("SO_HUU_NHA_NUOC_101_MA_2020_2025.xlsx"))
    args = parser.parse_args()
    if args.max_pages <= 0 or args.dpi <= 0:
        parser.error("--max-pages và --dpi phải dương")
    if args.ocr_missing or args.audit:
        if args.manifest.suffix.lower() == ".csv":
            manifest = pd.read_csv(args.manifest, encoding="utf-8-sig")
        else:
            manifest = pd.read_excel(args.manifest)
        if args.ocr_missing:
            ocr_statistics = ocr_missing_pages(manifest, args.cache_root, args.ocr_cache_root, args.pdf_root, args.tessdata, args.max_pages, args.dpi)
            print(json.dumps(ocr_statistics, ensure_ascii=False), flush=True)
        if args.audit:
            if not args.results.exists():
                details = pd.DataFrame()
            elif args.results.suffix.lower() == ".csv":
                details = pd.read_csv(args.results, encoding="utf-8-sig")
            else:
                details = pd.read_excel(args.results, sheet_name="Chi_tiet")
            audit_statistics = audit_candidates(manifest, details, args.cache_root, args.ocr_cache_root, args.audit_cache_root, args.fallback_audit_root, args.ocr_audit_root, args.pdf_root, args.audit_model)
            print(json.dumps(audit_statistics, ensure_ascii=False), flush=True)
    statistics = build_review(args.manifest, args.results, args.cache_root, args.ocr_cache_root, args.audit_cache_root, args.fallback_audit_root, args.ocr_audit_root, args.pdf_root, args.output)
    print(json.dumps({**statistics, "File kết quả": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
