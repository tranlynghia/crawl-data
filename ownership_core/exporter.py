"""
ownership_core.exporter
=======================
Xuất file Excel chuẩn mực duy nhất một sheet 'Co ty le so huu NN':
- Đủ 606 dòng dữ liệu tương ứng 101 mã x 6 năm (2020-2025).
- Không thêm sheet phụ; không lấy số năm khác điền vào mục thiếu.
- Định dạng chuyên nghiệp với openpyxl.
"""

from __future__ import annotations
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from typing import Any, Dict, List, Tuple
from ownership_core.inventory import InventoryItem
from ownership_core.validator import ValidationResult


EXCEL_HEADERS = [
    "STT",
    "Mã CK",
    "Tên công ty",
    "Năm Báo Cáo",
    "Ngày Chốt Danh Sách",
    "Tỷ Lệ Sở Hữu Nhà Nước (%)",
    "Số Lượng Cổ Phiếu",
    "Tổ Chức Đại Diện Vốn Nhà Nước",
    "Cơ Sở Xác Định",
    "Bằng Chứng Trích Xuất",
    "File Nguồn",
    "Trang",
    "Trạng Thái / Lý Do Chưa Xác Minh"
]


def export_single_sheet_excel(
    inventory: List[InventoryItem],
    results_map: Dict[Tuple[str, int], ValidationResult],
    output_path: Path
) -> bool:
    """
    Xuất file Excel đúng chuẩn:
    - 1 Sheet duy nhất: 'Co ty le so huu NN'.
    - 606 dòng dữ liệu (1 dòng mỗi mã-năm).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Co ty le so huu NN"
    ws.views.sheetView[0].showGridLines = True

    # 1. Định dạng Tiêu đề
    ws.append(EXCEL_HEADERS)
    ws.row_dimensions[1].height = 28

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for col_idx in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Fills & Fonts cho dữ liệu
    fill_state = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Xanh lá mềm
    fill_non_state = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")  # Xám nhạt
    fill_warning = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Vàng cảnh báo

    font_regular = Font(name="Segoe UI", size=10)
    font_bold = Font(name="Segoe UI", size=10, bold=True)

    # 2. Ghi chính xác 606 dòng dữ liệu theo thứ tự Inventory
    for r_idx, item in enumerate(inventory, 2):
        pair = (item.ticker, item.year)
        val_res = results_map.get(pair)

        if val_res:
            ratio = val_res.ratio
            shares = val_res.share_count
            holder = val_res.holder_name
            date_str = val_res.cutoff_date
            basis = val_res.basis
            evidence = val_res.evidence
            page_no = val_res.page_number or ""
            status_desc = val_res.status
            if val_res.notes:
                status_desc += f" ({val_res.notes})"
        else:
            ratio = None
            shares = None
            holder = "Chưa xử lý"
            date_str = "N/A"
            basis = "Chưa có kết quả"
            evidence = ""
            page_no = ""
            status_desc = "CHUA_XU_LY" if item.is_disk_available else "THIEU_FILE_TREN_DIA"

        row_data = [
            item.stt,
            item.ticker,
            item.company_name,
            item.year,
            date_str,
            ratio,
            shares,
            holder,
            basis,
            evidence,
            item.file_name or "",
            page_no,
            status_desc
        ]
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20

        # Định dạng ô
        is_has_state = (val_res and val_res.status == "DA_XAC_MINH_CO_SHNN" and ratio and ratio > 0)
        is_unverified = (val_res and not val_res.is_verified)

        for c_idx in range(1, len(EXCEL_HEADERS) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border

            if c_idx in (1, 4):  # STT, Năm
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 2:     # Mã CK
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_bold
            elif c_idx == 5:     # Ngày chốt
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 6:     # Tỷ lệ %
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if ratio is not None:
                    cell.number_format = '0.00"%"'
                if is_has_state:
                    cell.fill = fill_state
                    cell.font = font_bold
                elif val_res and val_res.status == "DA_XAC_MINH_KHONG_CO_SHNN":
                    cell.fill = fill_non_state
                elif is_unverified:
                    cell.fill = fill_warning
            elif c_idx == 7:     # Số lượng CP
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if shares is not None:
                    cell.number_format = "#,##0"
            elif c_idx in (3, 8, 9, 11, 13):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 10:    # Bằng chứng
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c_idx == 12:    # Trang
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Độ rộng cột
    col_widths = {
        1: 6,   # STT
        2: 10,  # Mã CK
        3: 35,  # Tên công ty
        4: 12,  # Năm Báo Cáo
        5: 16,  # Ngày Chốt
        6: 18,  # Tỷ Lệ Nhà Nước (%)
        7: 18,  # Số Lượng CP
        8: 32,  # Tổ Chức Đại Diện
        9: 38,  # Cơ Sở Xác Định
        10: 45, # Bằng Chứng
        11: 25, # File Nguồn
        12: 10, # Trang
        13: 30  # Trạng Thái / Lý Do
    }
    for c_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    ws.freeze_panes = "C2"

    try:
        wb.save(output_path)
        return True
    except PermissionError:
        # Nếu file bị khóa (ví dụ Excel đang mở), lưu ra file tạm thời
        fallback = output_path.with_name(f"{output_path.stem}_latest.xlsx")
        wb.save(fallback)
        return False
